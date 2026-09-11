"""Enriched Excel workbook output."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .insights import action_list, price_sensitivity, token_rollup, visibility_gaps
from .metrics import entity_summary

PCT_COLS = {
    "mkt_ctr", "own_ctr", "mkt_cart_rate", "own_cart_rate", "mkt_checkout_rate", "own_checkout_rate",
    "mkt_cvr", "own_cvr", "mkt_search_to_purchase", "own_search_to_purchase",
    "impr_share", "click_share", "cart_share", "purch_share", "share_drift",
    "price_gap_click", "price_gap_cart", "price_gap_purch",
    "own_ctr_lo", "own_ctr_hi", "own_cvr_lo", "own_cvr_hi",
    "mkt_fast_ship_click_share", "mkt_fast_ship_purch_share",
}
INDEX_COLS = {"ctr_index", "cart_index", "checkout_index", "cvr_index", "biggest_leak_index"}
INT_COLS = {
    "sqp_score", "sqp_volume", "impr_total", "impr_own", "clicks_total", "clicks_own",
    "carts_total", "carts_own", "purch_total", "purch_own", "queries", "headroom_purchases",
}

DATA_ORDER = [
    "entity_id", "scope", "period_label", "query", "query_type", "attributes", "competitor",
    "quadrant", "quadrant_strength", "biggest_leak", "recommended_action",
    "sqp_score", "sqp_volume",
    "impr_total", "impr_own", "impr_share",
    "clicks_total", "clicks_own", "click_share",
    "carts_total", "carts_own", "cart_share",
    "purch_total", "purch_own", "purch_share", "share_drift",
    "mkt_ctr", "own_ctr", "ctr_index", "ctr_confidence", "ctr_signif", "own_ctr_lo", "own_ctr_hi",
    "mkt_cart_rate", "own_cart_rate", "cart_index",
    "mkt_checkout_rate", "own_checkout_rate", "checkout_index",
    "mkt_cvr", "own_cvr", "cvr_index", "cvr_confidence", "cvr_signif", "own_cvr_lo", "own_cvr_hi",
    "lost_clicks", "lost_purchases", "headroom_purchases", "visibility_value",
    "price_click_mkt", "price_click_own", "price_gap_click", "price_position",
    "price_cart_mkt", "price_cart_own", "price_gap_cart",
    "price_purch_mkt", "price_purch_own", "price_gap_purch",
    "mkt_fast_ship_click_share", "mkt_fast_ship_purch_share",
    "mkt_search_to_purchase", "own_search_to_purchase",
    "clicks_rate_amz", "carts_rate_amz", "purch_rate_amz",
    "period_type", "period_end", "source",
]

HEADER_FILL = PatternFill("solid", fgColor="0F3D2A")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style(ws, df: pd.DataFrame) -> None:
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        letter = get_column_letter(i)
        width = 12
        if col in ("query", "why", "recommended_action"):
            width = 48
        elif col in ("attributes", "entity_id", "source"):
            width = 22
        ws.column_dimensions[letter].width = width
        n = len(df)
        if n == 0:
            continue
        rng = f"{letter}2:{letter}{n + 1}"
        if col in PCT_COLS:
            for r in range(2, n + 2):
                ws.cell(row=r, column=i).number_format = "0.00%"
        elif col in INDEX_COLS:
            for r in range(2, n + 2):
                ws.cell(row=r, column=i).number_format = "0.00x"
            ws.conditional_formatting.add(
                rng, ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                    mid_type="num", mid_value=1, mid_color="FFFFFF",
                                    end_type="num", end_value=2, end_color="63BE7B"))
        elif col in INT_COLS:
            for r in range(2, n + 2):
                ws.cell(row=r, column=i).number_format = "#,##0"
        elif col.startswith("price_") or col in ("lost_clicks", "lost_purchases", "visibility_value",
                                                 "purchases_at_stake"):
            for r in range(2, n + 2):
                ws.cell(row=r, column=i).number_format = "#,##0.00"
    ws.row_dimensions[1].height = 32


def _write(writer: pd.ExcelWriter, name: str, df: pd.DataFrame) -> None:
    df = df.copy()
    for c in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[c]):
            df[c] = df[c].dt.date
    df.to_excel(writer, sheet_name=name, index=False)
    _style(writer.sheets[name], df)


def write_report(d: pd.DataFrame, ctx: dict, out: str | Path, top_n: int = 25) -> Path:
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    cols = [c for c in DATA_ORDER if c in d.columns] + [c for c in d.columns if c not in DATA_ORDER]
    data = d[cols].sort_values(["entity_id", "period_label", "sqp_volume"], ascending=[True, True, False])

    summary = entity_summary(d)
    quad = (d.groupby(["entity_id", "period_label", "quadrant"]).agg(
        queries=("query", "size"), sqp_volume=("sqp_volume", "sum"), impr_own=("impr_own", "sum"),
        clicks_own=("clicks_own", "sum"), purch_own=("purch_own", "sum"),
        lost_purchases=("lost_purchases", "sum")).reset_index())
    qtype = (d.groupby(["entity_id", "period_label", "query_type"]).agg(
        queries=("query", "size"), sqp_volume=("sqp_volume", "sum"), impr_total=("impr_total", "sum"),
        impr_own=("impr_own", "sum"), clicks_total=("clicks_total", "sum"), clicks_own=("clicks_own", "sum"),
        purch_total=("purch_total", "sum"), purch_own=("purch_own", "sum")).reset_index())
    qtype["impr_share"] = qtype.impr_own / qtype.impr_total.replace(0, float("nan"))
    qtype["own_ctr"] = qtype.clicks_own / qtype.impr_own.replace(0, float("nan"))
    qtype["mkt_ctr"] = qtype.clicks_total / qtype.impr_total.replace(0, float("nan"))
    qtype["own_cvr"] = qtype.purch_own / qtype.clicks_own.replace(0, float("nan"))
    qtype["mkt_cvr"] = qtype.purch_total / qtype.clicks_total.replace(0, float("nan"))
    qtype["ctr_index"] = qtype.own_ctr / qtype.mkt_ctr
    qtype["cvr_index"] = qtype.own_cvr / qtype.mkt_cvr

    actions = pd.concat(
        [action_list(g, top_n) for _, g in d.groupby(["entity_id", "period_label"])], ignore_index=True)
    vis = pd.concat(
        [visibility_gaps(g, top_n) for _, g in d.groupby(["entity_id", "period_label"])], ignore_index=True)
    price = price_sensitivity(d)
    tok = pd.concat(
        [token_rollup(g).assign(entity_id=e, period_label=p).head(60)
         for (e, p), g in d.groupby(["entity_id", "period_label"])], ignore_index=True)
    tok = tok[["entity_id", "period_label"] + [c for c in tok.columns if c not in ("entity_id", "period_label")]]
    brands = pd.DataFrame({
        "kind": ["own_brand"] * len(ctx["own_brand_tokens"]) + ["competitor_brand"] * len(ctx["competitor_tokens"]),
        "token": ctx["own_brand_tokens"] + ctx["competitor_tokens"],
    })
    definitions = pd.DataFrame([
        ("Market CTR", "Clicks: Total Count / Impressions: Total Count (per impression)"),
        ("Brand/Own CTR", "Clicks: ASIN|Brand Count / Impressions: ASIN|Brand Count"),
        ("Market CVR", "Purchases: Total Count / Clicks: Total Count"),
        ("Brand/Own CVR", "Purchases: ASIN|Brand Count / Clicks: ASIN|Brand Count"),
        ("ctr_index / cvr_index", "own rate / market rate. >1 = beating the market"),
        ("share_drift", "purchase share - impression share. >0 = you convert above your share of voice"),
        ("lost_clicks", "(market CTR - own CTR) x own impressions, floored at 0"),
        ("lost_purchases", "(market CVR - own CVR) x own clicks, floored at 0"),
        ("headroom_purchases", "market purchases not made on your listing"),
        ("visibility_value", "extra purchases per +1pt impression share at market CTR x CVR"),
        ("price_gap_*", "(own median price - market median price) / market median price"),
        ("*_confidence", "sample size vs thresholds (own impressions for CTR, own clicks for CVR)"),
        ("*_signif", "Wilson 95% interval of own rate vs market rate: above / below / ns"),
        ("quadrant", "Winner / Leaky page / Hidden gem / Mismatch from ctr_index & cvr_index (>=1 vs <1)"),
        ("*_rate_amz", "Amazon's own rates: count / Search Query Volume (per search, can exceed 100%)"),
    ], columns=["metric", "definition"])

    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        _write(xw, "Summary", summary)
        _write(xw, "Actions", actions)
        _write(xw, "Quadrants", quad)
        _write(xw, "Query Types", qtype)
        _write(xw, "Visibility Gaps", vis)
        _write(xw, "Price Sensitivity", price)
        _write(xw, "Tokens", tok)
        _write(xw, "Data", data)
        _write(xw, "Brand Tokens", brands)
        _write(xw, "Definitions", definitions)
    return out
