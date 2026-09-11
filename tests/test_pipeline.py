"""End-to-end checks against the sample workbook (skipped if it is not present)."""
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from sqp_tool.insights import enrich
from sqp_tool.metrics import compute_metrics, wilson
from sqp_tool.parse import load_sqp

SAMPLE = Path.home() / "Desktop" / "SQP.xlsx"
pytestmark = pytest.mark.skipif(not SAMPLE.exists(), reason="sample workbook not present")


@pytest.fixture(scope="module")
def enriched():
    d = compute_metrics(load_sqp(SAMPLE))
    return enrich(d)


def test_parses_every_sheet_with_metadata():
    raw = load_sqp(SAMPLE)
    assert raw.entity_id.nunique() == 7
    assert set(raw.scope) == {"asin", "brand"}
    assert set(raw.period_label) >= {"2025 January", "2024 Q4", "2025 Q1", "2025 Q2"}
    assert raw["query"].str.len().min() > 0


def test_core_rates_match_manual_formulas():
    """Market CTR = K/H, Brand CTR = M/I, Market CVR = AC/K, Brand CVR = AE/M in the source sheet."""
    d = compute_metrics(load_sqp(SAMPLE, sheets=["AU_dm43"]))
    ws = openpyxl.load_workbook(SAMPLE, data_only=False)["AU_dm43"]
    hdr = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}
    for r in range(3, 13):
        q = str(ws.cell(r, 1).value).strip().lower()
        row = d[d["query"] == q].iloc[0]
        H, I, K, M, AC, AE = (ws.cell(r, hdr[h]).value or 0 for h in (
            "Impressions: Total Count", "Impressions: ASIN Count", "Clicks: Total Count",
            "Clicks: ASIN Count", "Purchases: Total Count", "Purchases: ASIN Count"))
        assert row.mkt_ctr == pytest.approx(K / H)
        assert (np.isnan(row.own_ctr) if I == 0 else row.own_ctr == pytest.approx(M / I))
        assert (np.isnan(row.mkt_cvr) if K == 0 else row.mkt_cvr == pytest.approx(AC / K))
        assert (np.isnan(row.own_cvr) if M == 0 else row.own_cvr == pytest.approx(AE / M))


def test_amazon_rates_are_per_search_volume():
    d = compute_metrics(load_sqp(SAMPLE, sheets=["AU_dm41"]))
    row = d[d["query"] == "avlt dual monitor arm"].iloc[0]
    assert row.clicks_rate_amz == pytest.approx(100 * row.clicks_total / row.sqp_volume, rel=1e-3)
    assert row.mkt_ctr < 0.1  # per-impression CTR is a real funnel rate


def test_wilson_bounds():
    lo, hi = wilson(pd.Series([0, 5, 10]), pd.Series([0, 10, 10]))
    assert np.isnan(lo[0]) and np.isnan(hi[0])
    assert 0 < lo[1] < 0.5 < hi[1] < 1
    assert hi[2] == pytest.approx(1.0)


def test_quadrants_and_brand_inference(enriched):
    d, ctx = enriched
    assert set(d.quadrant) <= {"Winner", "Leaky page", "Hidden gem", "Mismatch", "Not shown", "Insufficient data"}
    assert "avlt" in ctx["own_brand_tokens"]
    assert {"choco", "nose"} <= set(ctx["own_brand_tokens"])
    # a brand query is classified as own_brand
    assert (d[d["query"] == "avlt"].query_type == "own_brand").all()
    # low-sample rows never land in a judgement quadrant
    low = d[d.impr_own < 20]
    assert (low.quadrant == "Insufficient data").all()


def test_opportunity_never_negative(enriched):
    d, _ = enriched
    assert (d.lost_clicks >= 0).all() and (d.lost_purchases >= 0).all()
